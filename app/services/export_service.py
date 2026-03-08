import csv
import io
import logging

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models.lead import Lead

logger = logging.getLogger(__name__)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

EXPORT_CSV_COLUMNS = [
    "ig_username",
    "ig_full_name",
    "ig_bio_clean",
    "ig_website",
    "ig_category",
    "ig_follower_count",
    "score",
    "lead_category",
    "dm_message",
]


class ExportService:
    async def export_csv(self, campaign_id: str, db: AsyncSession) -> str:
        """Export DM-ready leads as CSV string for JarveePro import."""
        leads = await self._get_dm_ready_leads(campaign_id, db)

        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=EXPORT_CSV_COLUMNS)
        writer.writeheader()

        for lead in leads:
            writer.writerow({
                "ig_username": lead.ig_username,
                "ig_full_name": lead.ig_full_name or "",
                "ig_bio_clean": lead.ig_bio_clean or "",
                "ig_website": lead.ig_website or "",
                "ig_category": lead.ig_category or "",
                "ig_follower_count": lead.ig_follower_count or 0,
                "score": lead.score or 0,
                "lead_category": lead.lead_category or "",
                "dm_message": lead.dm_message or "",
            })

        return output.getvalue()

    async def export_json(self, campaign_id: str, db: AsyncSession) -> list[dict]:
        """Export DM-ready leads as JSON list."""
        leads = await self._get_dm_ready_leads(campaign_id, db)

        return [
            {
                "ig_username": lead.ig_username,
                "ig_full_name": lead.ig_full_name,
                "ig_bio_clean": lead.ig_bio_clean,
                "ig_website": lead.ig_website,
                "ig_category": lead.ig_category,
                "ig_follower_count": lead.ig_follower_count,
                "score": lead.score,
                "score_reason": lead.score_reason,
                "lead_category": lead.lead_category,
                "research_summary": lead.research_summary,
                "dm_message": lead.dm_message,
                "dm_variant_b": lead.dm_variant_b,
            }
            for lead in leads
        ]

    async def export_excel(self, campaign_id: str, db: AsyncSession) -> bytes:
        """Export DM-ready leads as Excel (.xlsx) with formatting."""
        if not HAS_OPENPYXL:
            raise ImportError("openpyxl is required for Excel export")

        leads = await self._get_dm_ready_leads(campaign_id, db)

        wb = Workbook()
        ws = wb.active
        ws.title = "DM Ready Leads"

        # Header style
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        headers = [
            "Username", "Full Name", "Bio", "Website", "IG Category",
            "Followers", "Score", "Category", "Research", "DM (Variant A)", "DM (Variant B)",
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        # Score color helper
        green_fill = PatternFill(start_color="00C853", end_color="00C853", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFD600", end_color="FFD600", fill_type="solid")

        for row_idx, lead in enumerate(leads, 2):
            ws.cell(row=row_idx, column=1, value=f"@{lead.ig_username}").border = thin_border
            ws.cell(row=row_idx, column=2, value=lead.ig_full_name or "").border = thin_border
            ws.cell(row=row_idx, column=3, value=lead.ig_bio_clean or "").border = thin_border
            ws.cell(row=row_idx, column=4, value=lead.ig_website or "").border = thin_border
            ws.cell(row=row_idx, column=5, value=lead.ig_category or "").border = thin_border
            ws.cell(row=row_idx, column=6, value=lead.ig_follower_count or 0).border = thin_border

            score_cell = ws.cell(row=row_idx, column=7, value=lead.score or 0)
            score_cell.border = thin_border
            score_cell.alignment = Alignment(horizontal="center")
            if lead.score and lead.score >= 80:
                score_cell.fill = green_fill
            elif lead.score and lead.score >= 70:
                score_cell.fill = yellow_fill

            ws.cell(row=row_idx, column=8, value=lead.lead_category or "").border = thin_border
            ws.cell(row=row_idx, column=9, value=lead.research_summary or "").border = thin_border
            ws.cell(row=row_idx, column=10, value=lead.dm_message or "").border = thin_border
            ws.cell(row=row_idx, column=11, value=lead.dm_variant_b or "").border = thin_border

        # Column widths
        widths = [18, 20, 35, 25, 15, 12, 8, 15, 40, 50, 50]
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

        # Auto-filter
        ws.auto_filter.ref = ws.dimensions

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    async def _get_dm_ready_leads(
        self, campaign_id: str, db: AsyncSession
    ) -> list[Lead]:
        result = await db.execute(
            select(Lead).where(
                Lead.campaign_id == campaign_id,
                Lead.status == "dm_ready",
                Lead.dm_message.isnot(None),
            ).order_by(Lead.score.desc())
        )
        return list(result.scalars().all())


export_service = ExportService()
